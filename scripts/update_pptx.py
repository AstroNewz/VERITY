import pptx
import os
import shutil

def update_shape(shape, new_text):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    if not tf.paragraphs:
        return
    
    first_p = tf.paragraphs[0]
    template_run = first_p.runs[0] if first_p.runs else None
    font_name = template_run.font.name if template_run else None
    font_size = template_run.font.size if template_run else None
    bold = template_run.font.bold if template_run else None
    italic = template_run.font.italic if template_run else None
    color_rgb = None
    if template_run and template_run.font.color and template_run.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB:
        color_rgb = template_run.font.color.rgb

    lines = new_text.split("\n")
    
    # Update first paragraph
    first_p.text = ""
    r0 = first_p.add_run()
    r0.text = lines[0]
    if font_name:
        r0.font.name = font_name
    if font_size:
        r0.font.size = font_size
    if bold is not None:
        r0.font.bold = bold
    if italic is not None:
        r0.font.italic = italic
    if color_rgb:
        r0.font.color.rgb = color_rgb
    
    # If there are existing extra paragraphs, reuse or clear them
    for p_idx in range(1, max(len(tf.paragraphs), len(lines))):
        if p_idx < len(lines):
            if p_idx < len(tf.paragraphs):
                p = tf.paragraphs[p_idx]
                p.text = ""
            else:
                p = tf.add_paragraph()
            r = p.add_run()
            r.text = lines[p_idx]
            if font_name:
                r.font.name = font_name
            if font_size:
                r.font.size = font_size
            if bold is not None:
                r.font.bold = bold
            if italic is not None:
                r.font.italic = italic
            if color_rgb:
                r.font.color.rgb = color_rgb
        else:
            if p_idx < len(tf.paragraphs):
                tf.paragraphs[p_idx].text = ""

def apply_deck_updates(input_pptx: str, output_pptx: str):
    prs = pptx.Presentation(input_pptx)
    
    # --- SLIDE 1 ---
    s1 = prs.slides[0].shapes
    update_shape(s1[6], "Round 3: Product Completion (Top 50 Finalist)   |   Live: cosmicboom.xyz & verity-brand-governance.vercel.app")
    update_shape(s1[13], "Rule C-12: no approved source (Brand Book pg. 14)")
    
    # --- SLIDE 2 ---
    s2 = prs.slides[1].shapes
    update_shape(s2[29], "ROUND 3 BUILD VALIDATION  Dual-set benchmark on 24 planted defects: 100% core recall, 0 false alarms.")
    
    # --- SLIDE 3 ---
    s3 = prs.slides[2].shapes
    update_shape(s3[37], "ZERO-ASSISTANCE JUDGE TOUR  In-situ guided tutorial with 4 persona tracks (Creator, Reviewer, Owner, All) & spotlight box.")
    
    # --- SLIDE 4 ---
    s4 = prs.slides[3].shapes
    update_shape(s4[4], "Upload / Draft")
    update_shape(s4[5], "Text, image alt-text, hex contrast, docx, pdf.")
    update_shape(s4[8], "Pre-Checks")
    update_shape(s4[9], "Deterministic regex checks casing & codenames first.")
    update_shape(s4[12], "Claim Ledger")
    update_shape(s4[13], "Verifies claims & 30-day expiry against evidence.")
    update_shape(s4[16], "Groq AI Voice")
    update_shape(s4[17], "Sub-second (<800ms) Llama 3.3 tone review via Groq.")
    update_shape(s4[20], "Intent Lock™")
    update_shape(s4[21], "Guarantees suggested fixes preserve marketing hook.")
    update_shape(s4[24], "Publish Passport")
    update_shape(s4[25], "Reviewer signs off; computes SHA-256 tamper seal.")
    update_shape(s4[26], "Steps 2 to 5 take <800ms. Step 6 computes an immutable SHA-256 seal. Humans always decide.")
    update_shape(s4[27], "Value: 100% pre-publication safety, cryptographic audit proof, and zero AI sterilization.")
    
    # --- SLIDE 5 ---
    s5 = prs.slides[4].shapes
    update_shape(s5[0], "Core features: Round 3 working build (Top 50 Finalist)")
    update_shape(s5[6], "Live rules with ON/OFF toggles, Brand Book citations, and session scoping.")
    update_shape(s5[11], "Claim Ledger™")
    update_shape(s5[12], "4-state lifecycle (Verified, Expires Soon <30d, Expired, Unsupported).")
    update_shape(s5[17], "Deterministic + AI")
    update_shape(s5[18], "Regex casing & codenames run first; Groq Llama 3.3 evaluates tone.")
    update_shape(s5[23], "Citing Findings")
    update_shape(s5[24], "Quotes exact substring, rule ID, and official Brand Book page citation.")
    update_shape(s5[29], "Reviewer Queue")
    update_shape(s5[30], "Side-by-side diffs, Intent Lock badges, and editorial feedback notes.")
    update_shape(s5[35], "Publish Passport")
    update_shape(s5[36], "Cryptographic SHA-256 seal with interactive real-time Tamper Simulator.")
    
    # --- SLIDE 6 ---
    s6 = prs.slides[5].shapes
    update_shape(s6[34], "Plus: In-situ self-guided tutorial with spotlight box & dual-link deployment (cosmicboom.xyz & Vercel mirror).")
    
    # --- SLIDE 7 ---
    s7 = prs.slides[6].shapes
    update_shape(s7[0], "What a reviewer & creator see")
    update_shape(s7[32], "The claim is flagged in context with Brand Book citation.")
    update_shape(s7[35], "Rule cited with intent-preserving fix (Intent Lock™).")
    update_shape(s7[38], "Reviewer signs off -> SHA-256 Publish Passport sealed.")
    
    # --- SLIDE 8 ---
    s8 = prs.slides[7].shapes
    update_shape(s8[0], "System architecture & 100% evaluation uptime")
    update_shape(s8[3], "Web Frontend")
    update_shape(s8[4], "Next.js 14 App Router (Prerendered, Neo-Brutalist, 360-412px Android Certified)")
    update_shape(s8[7], "Backend API")
    update_shape(s8[8], "FastAPI on Render Docker (24/7 cron keepalive, sub-second latency)")
    update_shape(s8[13], "Pipeline Engine")
    update_shape(s8[14], "Deterministic: Regex casing & secret codename guard\nPrivacy Gateway: Zero PII transit to external LLMs\nGroq Cloud SDK: Sub-second (<800ms) Llama 3.3 review\nIntent Lock: In-memory token restorer & hook preservation")
    update_shape(s8[17], "Evaluation Uptime")
    update_shape(s8[18], "Dual-Link Redundancy: Primary cosmicboom.xyz + Native Vercel mirror (instant fallback)")
    update_shape(s8[23], "Session Sandboxing: Multi-tenant org_{session_id} isolates judge workspaces")
    update_shape(s8[26], "LIVE VERIFIER: https://www.cosmicboom.xyz/verify/pass_sub_demo_1")
    update_shape(s8[31], "Zero PII Leakage")
    update_shape(s8[32], "Codenames & numbers tokenized before Groq AI.")
    update_shape(s8[33], "Cold-Start Alert")
    update_shape(s8[34], "On-screen banner notifies if cloud container is spinning up.")
    update_shape(s8[35], "Tamper Detection")
    update_shape(s8[36], "SHA-256 seal exposes DB mutations in real time.")
    
    # --- SLIDE 9 ---
    s9 = prs.slides[8].shapes
    update_shape(s9[0], "Round 2 Commitments Delivered 100% in Round 3")
    update_shape(s9[2], "ROUND 1: PLAN (DONE)")
    update_shape(s9[3], "Problem research\nUI wireframes\nSystem architecture")
    update_shape(s9[5], "ROUND 2: BUILD (DONE)")
    update_shape(s9[6], "Full working workflow\nGroq AI (<800ms)\nClaim Ledger & Expiry\nPublish Passport & Tamper\nSelf-Guided 4-Track Tour")
    update_shape(s9[8], "ROUND 3: 3 MAJOR ADDS (SHIPPED)")
    update_shape(s9[9], "1. Dense Vector Embeddings (pgvector claim matching)\n2. Multimodal Vision (poster, logo, color review)\n3. Figma & CMS Plugins (in-situ creator checks)")
    update_shape(s9[11], "BENCHMARK RESULTS")
    update_shape(s9[14], "100%")
    update_shape(s9[15], "Core Recall (24/24 caught)")
    update_shape(s9[18], "0")
    update_shape(s9[19], "False alarms (100% precision)")
    update_shape(s9[22], "12/12")
    update_shape(s9[23], "Passing Pytest unit tests")
    update_shape(s9[37], "Dual-Set Benchmark: Seeded set (12/12) + Fresh unseen set (12/12), tested via scripts/run_accuracy_check.py.")
    update_shape(s9[38], "Round 3 Additions: Substantial algorithmic & multimodal expansions, not cosmetic tweaks.")
    
    # --- SLIDE 10 ---
    s10 = prs.slides[9].shapes
    update_shape(s10[1], "Submission Links")
    update_shape(s10[3], "VERITY DELIVERABLES")
    update_shape(s10[8], "Primary: cosmicboom.xyz\nMirror: verity-brand-governance.vercel.app")
    update_shape(s10[13], "Backend: verity-api.onrender.com\nRepo: github.com/AstroNewz/VERITY")
    update_shape(s10[18], "Deliverables: PDF Report,\nInteractive Verifier & Tour Guide")
    update_shape(s10[25], "Team BraKet | CodeAxis Round 3: Product Completion | Top 50 Finalist")
    
    prs.save(output_pptx)
    print(f"Presentation saved to: {output_pptx}")

if __name__ == "__main__":
    root_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    
    # Backup original
    backup_file = os.path.join(root_dir, "Verity_Round1_Deck_Original_Backup.pptx")
    target_round1 = os.path.join(root_dir, "Verity_Round1_Deck_Complete.pptx")
    target_round2 = os.path.join(root_dir, "Verity_Round2_Deck_Complete.pptx")
    docs_round2 = os.path.join(root_dir, "docs", "Verity_Round2_Deck_Complete.pptx")
    
    if not os.path.exists(backup_file):
        shutil.copyfile(target_round1, backup_file)
        print(f"Backed up original to: {backup_file}")
        
    # Update target_round2
    apply_deck_updates(backup_file, target_round2)
    # Also update Verity_Round1_Deck_Complete.pptx directly as requested by the user
    apply_deck_updates(backup_file, target_round1)
    # Copy to docs/
    shutil.copyfile(target_round2, docs_round2)
    print("All PPTX files successfully updated!")
