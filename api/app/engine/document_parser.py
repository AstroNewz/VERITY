import io
from typing import Tuple

def parse_uploaded_file(filename: str, file_bytes: bytes) -> Tuple[str, str]:
    """
    Parses uploaded document files and extracts plain text.
    Supported extensions: .txt, .md, .docx, .pptx, .pdf
    Returns: (extracted_text, detected_type)
    """
    ext = filename.lower().split(".")[-1] if "." in filename else "txt"
    file_stream = io.BytesIO(file_bytes)

    if ext in ["txt", "md"]:
        try:
            return file_bytes.decode("utf-8"), "text"
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="ignore"), "text"

    elif ext == "docx":
        try:
            import docx
            doc = docx.Document(file_stream)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs), "docx"
        except Exception as e:
            return f"[Error parsing DOCX: {str(e)}]", "docx"

    elif ext == "pptx":
        try:
            import pptx
            prs = pptx.Presentation(file_stream)
            slide_texts = []
            for idx, slide in enumerate(prs.slides):
                texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            t = "".join([run.text for run in paragraph.runs if run.text.strip()])
                            if t:
                                texts.append(t)
                if texts:
                    slide_texts.append(f"--- Slide {idx + 1} ---\n" + "\n".join(texts))
            return "\n\n".join(slide_texts), "pptx"
        except Exception as e:
            return f"[Error parsing PPTX: {str(e)}]", "pptx"

    elif ext == "pdf":
        try:
            import pdfplumber
            with pdfplumber.open(file_stream) as pdf:
                page_texts = []
                for idx, page in enumerate(pdf.pages):
                    t = page.extract_text()
                    if t and t.strip():
                        page_texts.append(f"--- Page {idx + 1} ---\n" + t.strip())
                return "\n\n".join(page_texts), "pdf"
        except Exception as e:
            return f"[Error parsing PDF: {str(e)}]", "pdf"

    return file_bytes.decode("utf-8", errors="ignore"), "unknown"
