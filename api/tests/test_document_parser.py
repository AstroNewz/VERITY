import io
import pytest
from app.engine.document_parser import parse_uploaded_file

def test_parse_text_file():
    sample_text = "NovaFest is back with 10,000 attendees."
    text_bytes = sample_text.encode("utf-8")
    extracted, detected_type = parse_uploaded_file("announcement.txt", text_bytes)

    assert extracted == sample_text
    assert detected_type == "text"

def test_parse_markdown_file():
    sample_md = "# Nova Fest 2026\n\nTickets are affordable."
    md_bytes = sample_md.encode("utf-8")
    extracted, detected_type = parse_uploaded_file("readme.md", md_bytes)

    assert "Nova Fest 2026" in extracted
    assert detected_type == "text"

def test_parse_docx_file():
    import docx
    doc = docx.Document()
    doc.add_paragraph("Nova Fest confidential lineup announcement.")
    buf = io.BytesIO()
    doc.save(buf)
    extracted, detected_type = parse_uploaded_file("deck.docx", buf.getvalue())

    assert "Nova Fest confidential lineup announcement." in extracted
    assert detected_type == "docx"

def test_parse_pptx_file():
    import pptx
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Nova Fest Keynote 2026"
    buf = io.BytesIO()
    prs.save(buf)
    extracted, detected_type = parse_uploaded_file("presentation.pptx", buf.getvalue())

    assert "Nova Fest Keynote 2026" in extracted
    assert detected_type == "pptx"

